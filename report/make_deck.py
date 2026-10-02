"""Build the final presentation (report/LTH_Warmup_Final.pptx) in the mid-review deck's visual style.

Usage:  python -m analysis.plots && python -m analysis.final_figs && python report/make_deck.py

Numbers for the follow-up runs (F1, F2) are read from results/summary.csv, so the deck can be rebuilt as results change.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures" / "final"
OUT = ROOT / "report" / "LTH_Warmup_Final.pptx"

NAVY, NAVY2, PANEL = "14213D", "1F2E50", "F3F5F9"
ORANGE_D, ORANGE_L = "F2A65A", "B5651D"   # accent on dark / on light slides
WHITE, MUTED_D, SUB_D, MUTED_L = "FFFFFF", "9FB0CC", "C9D3E3", "5A6478"
SERIF, SANS = "Cambria", "Calibri"
TOTAL = 14


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0

    def slide(self, dark: bool):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
        self.n += 1
        return s

    @staticmethod
    def text(s, x, y, w, h, runs, size=15, color=NAVY, font=SANS, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=1.05, para_space=0):
        """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, overrides) tuples."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        paras = [runs] if isinstance(runs, str) else runs
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            p.space_after = Pt(para_space)
            for txt, ov in ([(para, {})] if isinstance(para, str) else para):
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name = ov.get("font", font)
                f.size = Pt(ov.get("size", size))
                f.bold = ov.get("bold", bold)
                f.color.rgb = rgb(ov.get("color", color))
        return tb

    @staticmethod
    def box(s, x, y, w, h, fill):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def header(self, s, kicker, title, dark=False):
        self.text(s, 0.6, 0.45, 12.13, 0.3, kicker, size=12, bold=True, color=ORANGE_D if dark else ORANGE_L)
        self.text(s, 0.6, 0.8, 12.13, 0.75, title, size=28, bold=True, font=SERIF, color=WHITE if dark else NAVY)

    def footer(self, s, note, dark=False):
        c = MUTED_D if dark else MUTED_L
        if note:
            self.text(s, 0.6, 6.85, 11.53, 0.35, note, size=11, color=c)
        self.text(s, 12.23, 6.85, 0.5, 0.35, str(self.n), size=11, color=c, align=PP_ALIGN.RIGHT)

    def card(self, s, x, y, w, h, title, body, dark_card=False, title_size=16, body_size=13, tag=None):
        self.box(s, x, y, w, h, NAVY if dark_card else PANEL)
        self.text(s, x + 0.25, y + 0.2, w - 0.5, 0.45, title, size=title_size, bold=True, font=SERIF,
                  color=WHITE if dark_card else NAVY)
        if tag:
            self.text(s, x + 0.25, y + 0.2, w - 0.5, 0.45, tag, size=11, bold=True,
                      color=ORANGE_D if dark_card else ORANGE_L, align=PP_ALIGN.RIGHT)
        self.text(s, x + 0.25, y + 0.72, w - 0.5, h - 0.9, body, size=body_size,
                  color="E3E8F0" if dark_card else NAVY, para_space=5)

    def figure(self, s, path, x, y, w=None, h=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)

    def table(self, s, x, y, w, rows, col_w, size=12, row_h=0.36, bold_cols=()):
        t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            t.rows[i].height = Inches(row_h)
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(NAVY if i == 0 else (WHITE if i % 2 else PANEL))
                c.margin_left = c.margin_right = Inches(0.08)
                c.margin_top = c.margin_bottom = Inches(0.03)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame
                tf.paragraphs[0].text = ""
                r = tf.paragraphs[0].add_run()
                r.text = str(val)
                r.font.name = SANS
                r.font.size = Pt(size)
                r.font.bold = i == 0 or j in bold_cols
                r.font.color.rgb = rgb(WHITE if i == 0 else NAVY)
                tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        return t


# ---------- numbers read from results (follow-ups) ----------

def adv(df, cond, rnd, seed):
    g = df[(df.condition == cond) & (df["round"] == rnd) & (df.seed == seed)]
    tk, sh = g[g.variant == "ticket"].acc, g[g.variant == "shuffle"].acc
    return None if tk.empty or sh.empty else 100 * (tk.iloc[0] - sh.iloc[0])


def acc(df, cond, rnd, seed, variant="ticket"):
    g = df[(df.condition == cond) & (df["round"] == rnd) & (df.seed == seed) & (df.variant == variant)].acc
    return None if g.empty else 100 * g.iloc[0]


def f(v, fmt="{:+.2f}"):
    return "pending" if v is None else fmt.format(v).replace("-", "−")


def build() -> None:
    df = pd.read_csv(ROOT / "results" / "summary.csv")
    d = Deck()
    R = [3, 6, 8]
    MASKS = ["34.5%", "12.0%", "6.0%"]

    # 1. Title
    s = d.slide(dark=True)
    d.text(s, 0.6, 0.6, 12.13, 0.35, "MACHINE LEARNING · COURSE PROJECT · FINAL PRESENTATION", size=13, bold=True, color=ORANGE_D)
    d.text(s, 0.6, 1.05, 11.2, 1.7, "Why Do Lottery Tickets Need Learning-Rate Warmup?", size=40, bold=True, font=SERIF, color=WHITE)
    d.text(s, 0.6, 2.75, 11.4, 1.0, "We tested sharpness and weight correlation as explanations. Neither holds: "
           "high learning rates break the mask that pruning selects, not the training of a good mask.", size=18, color=SUB_D)
    for i, (k, v) in enumerate([("SETTING", "ResNet-20 · CIFAR-10 · iterative magnitude pruning to 6% of weights · 2 seeds"),
                                ("DESIGN", "4 learning-rate conditions + a pre-registered causal test with the mask held fixed")]):
        x = 0.6 + i * 6.22
        d.box(s, x, 4.1, 5.92, 1.6, NAVY2)
        d.text(s, x + 0.3, 4.3, 5.32, 0.3, k, size=12, bold=True, color=ORANGE_D)
        d.text(s, x + 0.3, 4.65, 5.32, 1.0, v, size=15, color=WHITE)
    d.text(s, 0.6, 6.25, 12.13, 0.4, "Nitai S Koundinya  ·  Umair Ahmed Nawaz", size=16, bold=True, color=WHITE)
    d.text(s, 0.6, 6.65, 12.13, 0.35, "October 2026", size=12, color=MUTED_D)

    # 2. Question and hypotheses
    s = d.slide(dark=True)
    d.text(s, 0.6, 0.45, 12.13, 0.3, "RESEARCH QUESTION", size=12, bold=True, color=ORANGE_D)
    d.text(s, 0.6, 0.85, 12.13, 1.5, "Frankle & Carbin (2019): deep networks yield winning tickets only at a small learning rate or with warmup. "
           "Is that because high-LR training crosses the stability limit (sharpness), or because it decorrelates weights from θ₀ (Liu et al. 2021)?",
           size=22, font=SERIF, color=WHITE)
    hyps = [("H1 · Replicate", "Tickets win at η = 0.01 and with warmup, lose at η = 0.1"),
            ("H2 · Stability ratio", "Tickets fail where S = η·λmax / (2+2β) reaches ≈ 1 early on"),
            ("H3 · Warmup mechanism", "Warmup keeps S below 1 as the ticket advantage returns"),
            ("H4 · Causal test", "Fix good masks, train at η = 0.1: does SAM or an L2 anchor to θ₀ rescue them?")]
    for i, (h, b) in enumerate(hyps):
        x, y = 0.6 + (i % 2) * 6.22, 2.75 + (i // 2) * 1.75
        d.box(s, x, y, 5.92, 1.5, NAVY2)
        d.text(s, x + 0.3, y + 0.22, 5.32, 0.4, h, size=16, bold=True, color=ORANGE_D)
        d.text(s, x + 0.3, y + 0.65, 5.32, 0.75, b, size=15, color=WHITE)
    d.footer(s, "S = 1 is the edge of stability for SGD with heavy-ball momentum β = 0.9 (Cohen et al. 2021; Kalra & Barkeshli 2024).", dark=True)

    # 3. Method
    s = d.slide(dark=False)
    d.header(s, "METHOD", "Iterative magnitude pruning with a sharpness probe in every run")
    steps = [("1 · Init θ₀", "saved exactly"), ("2 · Train θ₀⊙m", "15k steps, log λmax, S"), ("3 · Evaluate", "acc, F1, AUC, R₀.₂"),
             ("4 · Prune 30%", "global magnitude"), ("5 · Reset to θ₀", "repeat ×8 → 6%")]
    for i, (a, b) in enumerate(steps):
        x = 0.6 + i * 2.53
        d.box(s, x, 1.8, 2.0, 1.05, NAVY)
        d.text(s, x + 0.15, 1.9, 1.75, 0.4, a, size=15, bold=True, color=WHITE)
        d.text(s, x + 0.15, 2.32, 1.75, 0.4, b, size=12, color=SUB_D)
        if i < 4:
            d.box(s, x + 2.12, 2.19, 0.3, 0.28, ORANGE_L)
    d.table(s, 0.6, 3.15, 12.13, [
        ["Condition", "η", "Warmup", "Role"],
        ["low", "0.01", "none", "tickets expected to win"],
        ["high", "0.1", "none", "tickets expected to fail"],
        ["warm03", "0.03", "linear, 10k steps", "Frankle's successful ResNet setting"],
        ["high_warm", "0.1", "linear, 10k steps", "same peak rate as high, with warmup"],
    ], [2.2, 1.2, 2.6, 6.13], size=13, row_h=0.38)
    d.text(s, 0.6, 5.2, 12.13, 0.9, [
        [("Baselines at every mask: ", {"bold": True}), ("random reinit and shuffled θ₀ (same layer-wise weights). ", {}),
         ("Ticket advantage", {"bold": True}), (" = ticket − best trained baseline. ", {}),
         ("Winning ticket", {"bold": True}), (" = advantage > 0 (2 seeds: > 2 SE) and accuracy within 0.5 pp of dense.", {})],
        [("Sharpness: ", {"bold": True}), ("λmax by 10-step Lanczos on 2,048 images, 15 points per run; H2 uses max S over steps 25–3,000. ", {}),
         ("Correlation: ", {"bold": True}), ("R₀.₂ = overlap of the top-20% weights of θ₀ and θ_T (chance ≈ 0.2).", {})],
    ], size=13, color=MUTED_L, para_space=6)
    d.footer(s, "ResNet-20 (0.27 M params) · SGD momentum 0.9 · batch 128 · LR ×0.1 at 10k and 12.5k · one T4 GPU, 3 runs in parallel under MPS.")

    # 4. H1
    s = d.slide(dark=False)
    d.header(s, "H1 · REPLICATION  —  SUPPORTED", "Tickets win everywhere except at η = 0.1 without warmup")
    d.figure(s, FIG / "advantage.png", 0.5, 1.7, w=8.1)
    d.card(s, 8.85, 1.75, 3.88, 2.35, "Two seeds agree", [
        "low: +2.1 → +6.3 pp", "warm03: +2.0 → +6.4 pp", "high: ≤ +0.5 pp, tickets 0.8–7.2 pp below dense: 0/4 winning"])
    d.card(s, 8.85, 4.3, 3.88, 2.3, "Warmup restores it", [
        "Same peak η = 0.1 with warmup (high_warm): +1.5 → +5.9 pp, 3/4 winning.",
        "Dense high is the most accurate network (89.7%), so this is not a weak baseline."], dark_card=True)
    d.footer(s, "Advantage = ticket − best trained baseline at the same mask. Bands: ± 1 SE over seeds. Gate A passed on both seeds.")

    # 5. Metrics table
    s = d.slide(dark=False)
    d.header(s, "RESULTS · TEST METRICS", "Dense networks and tickets at 34.5% and 6.0% of weights")
    mt = pd.read_csv(ROOT / "results" / "metrics_table.csv")
    rows = [["Condition", "Network", "Accuracy %", "Precision %", "Recall %", "F1 %", "AUC"]]
    for _, r in mt.iterrows():
        def m(c, scale=100, nd=2):
            return f"{scale * float(str(r[c]).split(' ')[0]):.{nd}f}"
        net = {"dense": "dense", "ticket 34.3%": "ticket 34.5%", "ticket 5.8%": "ticket 6.0%"}[r.network]
        a = m("acc") + (f" ± {100 * float(str(r['acc']).split('± ')[1]):.2f}" if "±" in str(r["acc"]) else "")
        rows.append([r.condition if net == "dense" else "", net, a, m("precision"), m("recall"), m("f1"), m("auc", 1, 4)])
    d.table(s, 0.6, 1.7, 12.13, rows, [1.8, 1.9, 2.0, 1.7, 1.6, 1.5, 1.63], size=12, row_h=0.36)
    d.footer(s, "Test set, 10,000 images · macro precision / recall / F1 · macro one-vs-rest ROC-AUC · mean ± SD over 2 seeds (high_warm: 1 seed).")

    # 6. H2
    s = d.slide(dark=False)
    d.header(s, "H2 · STABILITY RATIO  —  NOT SUPPORTED", "No ticket comes near the stability limit")
    d.figure(s, FIG / "max_s.png", 0.5, 1.7, w=7.9)
    d.card(s, 8.85, 1.75, 3.88, 2.55, "What fails", [
        "Max S over 105 runs is 0.54; failing tickets peak at 0.16–0.54.",
        "Within a condition S does not order the tickets (high: ρ = +0.26)."])
    d.card(s, 8.85, 4.5, 3.88, 2.1, "What holds", [
        "Only the failing condition has an early spike (max at step 25 in 15/18). Correlation, not yet cause."], dark_card=True)
    d.footer(s, "Caveat: with momentum the step-0 threshold is 2/λ (Kalra & Barkeshli 2024); dense high starts above it and still trains normally.")

    # 7. H3
    s = d.slide(dark=False)
    d.header(s, "H3 · WARMUP MECHANISM  —  PARTLY SUPPORTED", "Warmup removes the spike, but that is not what it fixes")
    d.card(s, 0.6, 1.75, 3.88, 3.6, "The coincidence holds", [
        "At η = 0.1, warmup brings the advantage back: 3/4 winning vs 0/4.",
        "It removes the step-25 spike: S ≤ 0.003 vs 0.13–0.54."], tag="✓", body_size=16)
    d.card(s, 4.73, 1.75, 3.88, 3.6, "S < 1 does not discriminate", [
        "high also stays below 1.",
        "By step 9,000 both runs sit on the same plateau: λmax 1.1–1.7, S 0.03–0.04.",
        "Warmup first lets λmax grow (25 → 53), as Kalra & Barkeshli predict."], tag="✗", body_size=16)
    d.card(s, 8.86, 1.75, 3.88, 3.6, "Confounded", [
        "Warmup also keeps more weight correlation (dense R₀.₂ 0.33 vs 0.25).",
        "It halves the mean LR over the first 10k steps.",
        "Only fixing the mask separates these. → H4"], dark_card=True, body_size=16)
    d.box(s, 0.6, 5.6, 12.13, 1.0, PANEL)
    d.text(s, 0.85, 5.6, 11.6, 1.0, [[("Warmup changes three things at once: ", {"bold": True}),
           ("the early sharpness spike, weight correlation and the mean learning rate. Phase A cannot say which one matters.", {})]],
           size=16, anchor=MSO_ANCHOR.MIDDLE)
    d.footer(s, "high_warm: seed 0 only.")

    # 8. H4 design
    s = d.slide(dark=False)
    d.header(s, "H4 · CAUSAL TEST  —  DESIGN", "Hold warmup's good masks fixed and change only the training")
    for i, (t, b) in enumerate([
        ("(a) plain", ["η = 0.1, no warmup", "Does a good mask still win under high-LR training?"]),
        ("(b) SAM, ρ = 0.05", ["Lowers sharpness, leaves correlation", "Rescue here → sharpness story"]),
        ("(c) L2 anchor to θ₀", ["(λ/2)·‖m⊙(θ − θ₀)‖², λ = 0.01 (chosen to match low's R₀.₂)", "Rescue here → correlation story"])]):
        d.card(s, 0.6 + i * 4.13, 1.75, 3.88, 2.4, t, b)
    d.box(s, 0.6, 4.4, 12.13, 1.75, NAVY)
    d.text(s, 0.9, 4.6, 11.5, 0.4, "PRE-REGISTERED RULE (committed before any H4 result)", size=13, bold=True, color=ORANGE_D)
    d.text(s, 0.9, 5.0, 11.5, 1.1, [
        "Rescue = ticket − shuffle ≥ +2.0 pp at ≥ 2 of 3 masks, and the intervention must move its target variable vs (a).",
        "Masks: warm03 seed 0 at 34.5%, 12.0%, 6.0%. The failing high chain never exceeds +0.7 pp, so the bar is well above failure."],
        size=15, color=WHITE, para_space=6)
    d.footer(s, "Each mask trained as ticket (θ₀⊙m) and shuffle (θ₀ permuted within each layer over the mask): 18 trainings, SAM at 2× cost.")

    # 9. H4 result
    s = d.slide(dark=False)
    d.header(s, "H4 · RESULT", "The good masks already win at η = 0.1. There is nothing to rescue")
    d.figure(s, FIG / "h4_interventions.png", 0.5, 1.7, w=7.9)
    d.card(s, 8.85, 1.75, 3.88, 1.6, "(a) plain", ["+3.6 and +4.8 pp: 2/3 over the bar. Seed 1: +3.5, +4.4."], tag="WINS")
    d.card(s, 8.85, 3.5, 3.88, 1.45, "(b) SAM", ["S lowered at 3/3 masks; advantage moves ≤ 0.5 pp."], tag="NO EFFECT")
    d.card(s, 8.85, 5.1, 3.88, 1.5, "(c) anchor", ["Ticket −3 pp, shuffle −10 pp vs (a). Not a better ticket."], dark_card=True, tag="HURTS")
    has_f1 = (df.condition.eq("h4_high") & df.seed.eq(1)).any()
    d.footer(s, "Seed 0. Ticket accuracy, (a): 89.27 / 88.91 / 87.24%." + (" Open circles: (a) on seed 1 (F1)." if has_f1 else ""))

    # 10. Mask decides
    s = d.slide(dark=False)
    d.header(s, "H4 · INTERPRETATION", "Same training, different mask: the mask decides")
    d.figure(s, FIG / "h4_masks.png", 0.5, 1.7, w=7.9)
    d.card(s, 8.85, 1.75, 3.88, 2.45, "Under identical training", [
        "η = 0.1, no warmup, same θ₀: warm03's 6% mask 87.2%, high's own 6% mask 82.7%.",
        "Shuffled controls equal (82.4 vs 82.2%)."])
    d.card(s, 8.85, 4.35, 3.88, 2.25, "Both stories fail", [
        "Good mask: spike present, R₀.₂ at chance, ticket wins.",
        "Bad mask + warmup: spike gone, ticket still fails (F2)."], dark_card=True)
    d.footer(s, "high and warm03 seed 0 share θ₀ bit for bit: the green and orange lines differ only in the mask, the orange and violet only in warmup.")

    # 11. Follow-ups
    s = d.slide(dark=False)
    d.header(s, "FOLLOW-UPS · PRE-REGISTERED", "Does it replicate, and does warmup rescue a bad mask?")
    f1 = [adv(df, "h4_high", r, 1) for r in R]
    f2 = [adv(df, "h4_swap_high_warm", r, 0) for r in R]
    n1, n2 = sum(v is not None and v >= 2 for v in f1), sum(v is not None and v >= 2 for v in f2)
    done1, done2 = sum(v is not None for v in f1), sum(v is not None for v in f2)
    v1 = "pending" if done1 < 3 else ("REPLICATES" if n1 >= 2 else "DOES NOT REPLICATE")
    v2 = "pending" if done2 < 3 else ("MASK FAILS" if n2 < 2 else "WARMUP RESCUES")
    rows = [["", "34.5%", "12.0%", "6.0%", "≥ +2 pp", "Verdict"],
            ["(a) seed 0  · warm03 mask, η 0.1", *[f(adv(df, "h4_high", r, 0)) for r in R], "2/3", "wins"],
            ["F1 · (a) on seed 1 masks + θ₀", *[f(v) for v in f1], f"{n1}/{done1}", v1.lower()],
            ["high chain · own masks, η 0.1", "—", "+0.39*", "+0.57", "0/2", "fails"],
            ["F2 · high's masks + warmup", *[f(v) for v in f2], f"{n2}/{done2}", v2.lower()]]
    d.table(s, 0.6, 1.75, 12.13, rows, [3.9, 1.5, 1.5, 1.5, 1.4, 2.33], size=13, row_h=0.45)
    t1 = [acc(df, "h4_swap_high_warm", r, 0) for r in R]
    t0 = [acc(df, "high", r, 0) for r in R]
    d.card(s, 0.6, 4.3, 5.92, 2.3, "F1 · replication", [
        "Same experiment as (a) on an independent seed: warm03 seed 1's masks and θ₀.",
        f"Advantage {' / '.join(f(v) for v in f1)} pp → {v1}."])
    d.card(s, 6.82, 4.3, 5.92, 2.3, "F2 · reverse swap", [
        "high's own masks, trained with warmup at η = 0.1 (high_warm schedule).",
        f"Ticket {' / '.join(f(v, '{:.2f}') for v in t1)}% vs {' / '.join(f(v, '{:.2f}') for v in t0)}% without warmup → {v2}."],
        dark_card=True)
    d.footer(s, "Ticket − shuffle (pp). *high chain at 12.0%: vs reinit (no shuffle run there). Rule as in H4, fixed before the runs started.")

    # 12. Verdicts
    s = d.slide(dark=True)
    d.text(s, 0.6, 0.45, 12.13, 0.3, "VERDICTS", size=12, bold=True, color=ORANGE_D)
    d.text(s, 0.6, 0.85, 12.13, 0.9, "High learning rates break the mask, not the training", size=30, bold=True, font=SERIF, color=WHITE)
    verdicts = [("H1 · Replicate", "SUPPORTED", "2 seeds: +2 to +6.4 pp at η 0.01 / warmup; ≤ +0.5 pp at η 0.1"),
                ("H2 · Stability ratio", "NOT SUPPORTED", "max S 0.54 ≪ 1; the early spike is not causal (H4)"),
                ("H3 · Warmup mechanism", "PARTLY", "Coincidence holds, but warmup is not needed to train a good mask and cannot fix a bad one"),
                ("H4 · Causal test", "MASK, NOT DYNAMICS", "Good masks win at η 0.1 (2 seeds); bad masks fail even with warmup")]
    for i, (h, tag, b) in enumerate(verdicts):
        x, y = 0.6 + (i % 2) * 6.22, 2.0 + (i // 2) * 2.05
        d.box(s, x, y, 5.92, 1.8, NAVY2)
        d.text(s, x + 0.3, y + 0.22, 5.32, 0.4, h, size=17, bold=True, color=WHITE)
        d.text(s, x + 0.3, y + 0.25, 5.32, 0.4, tag, size=12, bold=True, color=ORANGE_D, align=PP_ALIGN.RIGHT)
        d.text(s, x + 0.3, y + 0.75, 5.32, 0.9, b, size=15, color=SUB_D)
    d.footer(s, "Fits Paul et al. (2023): an IMP mask encodes the training run that produced it. Warmup matters while the mask is being found.", dark=True)

    # 13. Limitations and next steps
    s = d.slide(dark=False)
    d.header(s, "LIMITATIONS & NEXT STEPS", "What this study cannot say yet")
    d.card(s, 0.6, 1.75, 5.92, 4.1, "Limitations", body_size=16, title_size=18, body=[
        "Seeds: 2 for the main conditions and H4(a); 1 for high_warm, SAM, the anchor and the reverse swap.",
        "Good masks at η = 0.1 still sit 0.6–2.6 pp below dense: some LR cost survives.",
        "Shortened protocol: 15k steps, 30% per round, no rewinding, one model and dataset.",
        "Sharpness at 15 points per run; the minibatch stability threshold was not measured.",
        "One strength each for SAM (ρ 0.05) and the anchor (λ 0.01)."])
    d.card(s, 6.82, 1.75, 5.92, 4.1, "Next steps", body_size=16, title_size=18, dark_card=True, body=[
        "Switch η mid-chain to find the IMP round where masks go bad.",
        "Compare high vs warm03 masks directly: layer-wise density, overlap.",
        "Learning-rate rewinding (Frankle et al. 2020), known to restore tickets at high η: does it work by fixing the mask?"])
    d.footer(s, "")

    # 14. Close
    s = d.slide(dark=True)
    d.text(s, 0.6, 0.6, 12.13, 0.35, "CONCLUSION", size=13, bold=True, color=ORANGE_D)
    d.text(s, 0.6, 1.1, 12.13, 2.4, "Lottery tickets fail at high learning rates because iterative magnitude pruning picks worse masks, "
           "not because training becomes unstable or loses its correlation with θ₀.", size=30, bold=True, font=SERIF, color=WHITE)
    d.text(s, 0.6, 3.9, 12.13, 1.4, ["A pre-registered causal test, with the mask held fixed, separated three explanations that the observational data could not.",
                                      "Code, configs, every run and both write-ups: github.com/kneat2448/LTH-ML-final"], size=17, color=SUB_D, para_space=8)
    d.text(s, 0.6, 6.25, 12.13, 0.4, "Nitai S Koundinya  ·  Umair Ahmed Nawaz", size=16, bold=True, color=WHITE)
    d.text(s, 0.6, 6.65, 12.13, 0.35, "Thank you · questions welcome", size=12, color=MUTED_D)

    assert d.n == TOTAL, d.n
    d.prs.save(OUT)
    print(f"deck: {OUT} ({d.n} slides)")


if __name__ == "__main__":
    build()
